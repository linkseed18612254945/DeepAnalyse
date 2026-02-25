"""Intelligence Report Generator.

Produces structured analyst-grade reports in Markdown, PDF, DOCX, and JSON
formats from an AnalysisSession.
"""

from __future__ import annotations

import logging
from pathlib import Path

from ..core.config import settings
from ..core.llm import chat_completion
from ..core.models import (
    AnalysisResult,
    AnalysisSession,
    EventFraming,
    IntelligenceReport,
    RetrievalRoundResult,
)

logger = logging.getLogger(__name__)

REPORT_SYSTEM_PROMPT = """\
You are a professional intelligence report writer.  Given structured analytical
findings, produce a polished intelligence report following this exact format.

Write in clear, authoritative prose suitable for senior decision-makers.
Use the section headings exactly as given.  Be thorough but concise — each
section should be substantive (2-5 paragraphs), not padded.

Do NOT include a title line — the system will prepend that.

Sections to produce (use ## for each heading):

## Executive Summary
(2-3 paragraphs — standalone readable)

## Event Overview & Context
(Timeline, geography, verified facts)

## Key Actors Analysis
(Profiles, roles, relationships, motives)

## Deep Background
(Historical context, structural forces, precedents)

## Hidden Connections & Network Analysis
(Below-the-surface layer — indirect ties, financial links, etc.)

## Analytical Assessment
(What it means, why it matters, what's likely next)

## Scenarios & Implications
(Forward-looking with probability weighting, use bullet points for each scenario)

## Confidence Assessment
(What we know well, what is uncertain, what is unknown)

## Key Questions Remaining
(Gaps for further investigation, as a numbered list)

## Sources & Citations
(Categorised by type and reliability)

Guidelines:
- Every factual claim should reference a source where possible.
- Use [Source Title](URL) markdown link format for citations.
- Label confidence explicitly: (High Confidence), (Medium Confidence), etc.
- Be analytical, not merely descriptive.
"""


async def generate_report(session: AnalysisSession) -> IntelligenceReport:
    """Generate the full intelligence report from session data."""

    logger.info("Generating intelligence report for session %s", session.id)

    # Build the context prompt from session data
    context = _build_report_context(session)

    report_body = await chat_completion(
        system_prompt=REPORT_SYSTEM_PROMPT,
        user_prompt=context,
        max_tokens=8000,
        temperature=0.3,
    )

    # Parse sections from the markdown
    sections = _parse_sections(report_body)

    report = IntelligenceReport(
        title=f"Intelligence Report: {session.framing.core_event}" if session.framing else "Intelligence Report",
        executive_summary=sections.get("Executive Summary", ""),
        event_overview=sections.get("Event Overview & Context", ""),
        key_actors_analysis=sections.get("Key Actors Analysis", ""),
        deep_background=sections.get("Deep Background", ""),
        hidden_connections=sections.get("Hidden Connections & Network Analysis", ""),
        analytical_assessment=sections.get("Analytical Assessment", ""),
        scenarios_and_implications=sections.get("Scenarios & Implications", ""),
        confidence_assessment=sections.get("Confidence Assessment", ""),
        key_questions_remaining=sections.get("Key Questions Remaining", ""),
        sources_and_citations=sections.get("Sources & Citations", ""),
    )

    return report


def render_markdown(report: IntelligenceReport) -> str:
    """Render the report as a complete Markdown document."""

    parts = [
        f"# {report.title}",
        f"*Generated: {report.generated_at.strftime('%Y-%m-%d %H:%M UTC')}*\n",
        f"## Executive Summary\n{report.executive_summary}\n",
        f"## Event Overview & Context\n{report.event_overview}\n",
        f"## Key Actors Analysis\n{report.key_actors_analysis}\n",
        f"## Deep Background\n{report.deep_background}\n",
        f"## Hidden Connections & Network Analysis\n{report.hidden_connections}\n",
        f"## Analytical Assessment\n{report.analytical_assessment}\n",
        f"## Scenarios & Implications\n{report.scenarios_and_implications}\n",
        f"## Confidence Assessment\n{report.confidence_assessment}\n",
        f"## Key Questions Remaining\n{report.key_questions_remaining}\n",
        f"## Sources & Citations\n{report.sources_and_citations}\n",
    ]
    return "\n".join(parts)


async def save_report(
    report: IntelligenceReport,
    session_id: str,
    formats: list[str] | None = None,
) -> dict[str, str]:
    """Save the report to disk in the requested formats.

    Returns a dict of format -> file path.
    """

    output_dir = settings.reports_output_path / session_id
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, str] = {}

    target_formats = formats or ["markdown", "json"]
    md_content = render_markdown(report)

    if "markdown" in target_formats:
        md_path = output_dir / "report.md"
        md_path.write_text(md_content, encoding="utf-8")
        paths["markdown"] = str(md_path)

    if "json" in target_formats:
        import json
        json_path = output_dir / "report.json"
        json_path.write_text(
            json.dumps(report.model_dump(), indent=2, default=str),
            encoding="utf-8",
        )
        paths["json"] = str(json_path)

    if "pdf" in target_formats:
        try:
            pdf_path = output_dir / "report.pdf"
            _markdown_to_pdf(md_content, str(pdf_path))
            paths["pdf"] = str(pdf_path)
        except Exception as e:
            logger.warning("PDF generation failed: %s", e)

    if "docx" in target_formats:
        try:
            docx_path = output_dir / "report.docx"
            _markdown_to_docx(md_content, report.title, str(docx_path))
            paths["docx"] = str(docx_path)
        except Exception as e:
            logger.warning("DOCX generation failed: %s", e)

    return paths


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _build_report_context(session: AnalysisSession) -> str:
    parts: list[str] = []

    if session.framing:
        f = session.framing
        parts.append(f"Event: {f.core_event}")
        parts.append(f"What: {f.what}")
        parts.append(f"Who: {', '.join(f.who)}")
        parts.append(f"Where: {', '.join(f.where)}")
        parts.append(f"When: {f.when}")
        parts.append(f"Why: {f.why}")
        parts.append(f"How: {f.how}")
        parts.append(f"Significance: {f.so_what}\n")

    for rr in session.retrieval_rounds:
        parts.append(f"--- Round {rr.round_number}: {rr.focus} ---")
        parts.append(rr.findings_summary)
        sources_text = "\n".join(
            f"  - [{s.title}]({s.url}) (Tier {s.tier})" for s in rr.sources_found[:10]
        )
        parts.append(f"Sources:\n{sources_text}\n")

    if session.analysis:
        a = session.analysis
        parts.append("--- Analysis Results ---")
        for ap in a.actor_profiles:
            parts.append(f"Actor: {ap.name} — {ap.role} (Confidence: {ap.confidence})")
            parts.append(f"  Motivations: {', '.join(ap.motivations)}")
        for cc in a.causal_chains:
            parts.append(f"Causal: {cc.cause} → {cc.effect} ({cc.confidence})")
        if a.hidden_connections:
            parts.append(f"Hidden connections: {'; '.join(a.hidden_connections)}")
        if a.contradictions:
            parts.append(f"Contradictions: {'; '.join(a.contradictions)}")
        for s in a.scenarios:
            parts.append(f"Scenario: {s.title} ({s.probability}) — {s.description}")
        parts.append(f"Confidence: {a.confidence_assessment}")
        parts.append(f"Open questions: {'; '.join(a.key_questions_remaining)}")

    return "\n".join(parts)


def _parse_sections(markdown: str) -> dict[str, str]:
    """Split markdown by ## headings into a heading -> content dict."""
    sections: dict[str, str] = {}
    current_heading = ""
    current_lines: list[str] = []

    for line in markdown.split("\n"):
        if line.startswith("## "):
            if current_heading:
                sections[current_heading] = "\n".join(current_lines).strip()
            current_heading = line[3:].strip()
            current_lines = []
        else:
            current_lines.append(line)

    if current_heading:
        sections[current_heading] = "\n".join(current_lines).strip()

    return sections


def _markdown_to_pdf(md_content: str, output_path: str) -> None:
    """Convert markdown to PDF using weasyprint."""
    import markdown as md_lib
    from weasyprint import HTML

    html_content = md_lib.markdown(md_content, extensions=["tables", "fenced_code"])
    styled_html = f"""
    <html><head><style>
        body {{ font-family: 'Georgia', serif; margin: 2cm; line-height: 1.6; }}
        h1 {{ color: #1a1a2e; border-bottom: 2px solid #16213e; }}
        h2 {{ color: #16213e; margin-top: 1.5em; }}
        blockquote {{ border-left: 3px solid #0f3460; padding-left: 1em; color: #555; }}
    </style></head><body>{html_content}</body></html>
    """
    HTML(string=styled_html).write_pdf(output_path)


def _markdown_to_docx(md_content: str, title: str, output_path: str) -> None:
    """Convert markdown to DOCX using python-docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=0)

    for line in md_content.split("\n"):
        if line.startswith("# "):
            continue  # Skip the title (already added)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)

    doc.save(output_path)
